"""Build the casualty training table.

Sources: USGS ComCat PAGER exposure (2008 on), NOAA NCEI HazEL labels, USGS EXPO-CAT and PAGER-CAT
(1960-2007), World Bank income classes and GDP per capita, Natural Earth country outlines.

Everything downloaded is cached under data/raw, so a rerun only fetches what is missing.
Output: ml/casualty/artifacts/training_table.csv.gz and dataset_summary.json.
"""

import argparse
import json
import math
import threading
import time
import xml.etree.ElementTree as ET
import zipfile
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, timezone
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
import requests
from matplotlib.path import Path as MplPath

HERE = Path(__file__).resolve().parent
REPO = HERE.parents[1]
RAW = REPO / "data" / "raw"
CATALOGS = RAW / "catalogs"
PAGER_DIR = RAW / "pager"
ARTIFACTS = HERE / "artifacts"

USER_AGENT = "sentinel-er (+https://github.com/shivang-goliyan/sentinel-er)"
COMCAT = "https://earthquake.usgs.gov/fdsnws/event/1/query"
HAZEL = "https://www.ngdc.noaa.gov/hazel/hazard-service/api/v1/earthquakes"
EXPOCAT_URL = "https://earthquake.usgs.gov/static/lfs/data/pager/catalogs/EXPO_CAT_2007_12.csv"
PAGERCAT_URL = "https://www.sciencebase.gov/catalog/file/get/5bc730dde4b0fc368ebcad8a"
FATALITY_URL = "https://raw.githubusercontent.com/usgs/pager/master/losspager/data/fatality.xml"
OGHIST_URL = "https://datacatalogfiles.worldbank.org/ddh-published/0037712/DR0090754/OGHIST.xlsx"
NE_URL = ("https://raw.githubusercontent.com/nvkelso/natural-earth-vector/master/geojson/"
          "ne_50m_admin_0_countries.geojson")
WB_COUNTRIES = "https://api.worldbank.org/v2/country/all?format=json&per_page=400"
WB_GDP = ("https://api.worldbank.org/v2/country/all/indicator/NY.GDP.PCAP.CD"
          "?format=json&per_page=20000&date=1960:2025")

MMI_BINS = list(range(1, 11))
FEATURE_BINS = list(range(4, 11))
INCOME_ORDER = {"L": 0, "LM": 1, "UM": 2, "H": 3}
API_INCOME = {"LIC": "L", "LMC": "LM", "UMC": "UM", "HIC": "H"}

HOLDOUTS = {
    "turkey-2023": {"id": "us6000jllz", "time": "2023-02-06T01:17:34Z", "lat": 37.2256, "lon": 37.0143,
                    "extra_ids": ["us6000jlqa"]},
    "nepal-2015": {"id": "us20002926", "time": "2015-04-25T06:11:25Z", "lat": 28.2305, "lon": 84.7314,
                   "extra_ids": []},
}
HOLDOUT_RADIUS_KM = 200
HOLDOUT_BEFORE_DAYS = 30
HOLDOUT_AFTER_DAYS = 365

JOIN_SECONDS = 60
JOIN_KM = 100
JOIN_MAG = 0.3


class Fetcher:
    """requests with a global rate limit and retries."""

    def __init__(self, per_second=4.0):
        self.gap = 1.0 / per_second
        self.lock = threading.Lock()
        self.last = 0.0
        self.session = requests.Session()
        self.session.headers["User-Agent"] = USER_AGENT
        self.calls = 0

    def get(self, url, params=None, timeout=60):
        for attempt in range(6):
            with self.lock:
                wait = self.last + self.gap - time.monotonic()
                if wait > 0:
                    time.sleep(wait)
                self.last = time.monotonic()
                self.calls += 1
            try:
                r = self.session.get(url, params=params, timeout=timeout)
            except requests.RequestException:
                time.sleep(2 ** attempt)
                continue
            if r.status_code == 200:
                return r
            if r.status_code in (404, 204):
                return r
            if r.status_code == 400:
                r.raise_for_status()
            time.sleep(2 ** attempt + (5 if r.status_code == 429 else 0))
        raise RuntimeError(f"gave up on {url} {params}")


fetch = Fetcher()


def cached_file(url, path, params=None):
    path = Path(path)
    if not path.exists():
        path.parent.mkdir(parents=True, exist_ok=True)
        r = fetch.get(url, params=params, timeout=180)
        r.raise_for_status()
        path.write_bytes(r.content)
    return path


EPOCH = pd.Timestamp(0, tz="UTC")


def epoch_seconds(values):
    """Seconds since 1970 for a datetime Series or a Timestamp, whatever unit pandas stored it in."""
    if isinstance(values, pd.Timestamp):
        return (values - EPOCH).total_seconds()
    return (values - EPOCH).dt.total_seconds().to_numpy()


def haversine_km(lat1, lon1, lat2, lon2):
    lat1, lon1, lat2, lon2 = map(np.radians, (lat1, lon1, lat2, lon2))
    a = np.sin((lat2 - lat1) / 2) ** 2 + np.cos(lat1) * np.cos(lat2) * np.sin((lon2 - lon1) / 2) ** 2
    return 6371.0 * 2 * np.arcsin(np.sqrt(np.clip(a, 0, 1)))


# --- ComCat PAGER -----------------------------------------------------------------------------

def list_pager_events(refresh_year=None):
    spans = [("1900-01-01", "2008-01-01", "pre2008")]
    this_year = datetime.now(timezone.utc).year
    spans += [(f"{y}-01-01", f"{y + 1}-01-01", str(y)) for y in range(2008, this_year + 1)]
    rows = []
    for start, end, tag in spans:
        path = PAGER_DIR / "lists" / f"{tag}.geojson"
        if tag == str(refresh_year) and path.exists():
            path.unlink()
        cached_file(COMCAT, path, params={"producttype": "losspager", "starttime": start,
                                          "endtime": end, "format": "geojson", "limit": 20000})
        data = json.loads(path.read_text())
        if len(data["features"]) >= 20000:
            raise RuntimeError(f"{tag} hit the 20000 row limit, split it")
        for f in data["features"]:
            p = f["properties"]
            lon, lat, depth = f["geometry"]["coordinates"]
            rows.append({"event_id": f["id"], "time": pd.to_datetime(p["time"], unit="ms", utc=True),
                         "lat": lat, "lon": lon, "depth_km": depth, "magnitude": p["mag"],
                         "summary_alert": p.get("alert"), "place": p.get("place")})
    return pd.DataFrame(rows).drop_duplicates("event_id")


def pick_events(events, green_sample, seed):
    alert = events["summary_alert"].fillna("none")
    always = alert.isin(["yellow", "orange", "red"]) | (events["magnitude"] >= 6.0)
    holdout_ids = {h["id"] for h in HOLDOUTS.values()} | {i for h in HOLDOUTS.values() for i in h["extra_ids"]}
    always |= events["event_id"].isin(holdout_ids)
    pool = events[~always]
    n = min(green_sample, len(pool))
    sampled = pool.sample(n=n, random_state=seed)
    picked = pd.concat([events[always].assign(weight=1.0, sampled=False),
                        sampled.assign(weight=len(pool) / max(n, 1), sampled=True)])
    rule = {
        "always_kept": "PAGER alert yellow/orange/red, or magnitude >= 6.0, or a holdout event",
        "sampled": f"{n} of the remaining {len(pool)} green / no-alert events below M6.0, "
                   f"uniform random (seed {seed}), weight {len(pool) / max(n, 1):.3f} each",
        "always_kept_count": int(always.sum()),
        "listed_events": int(len(events)),
        "pool_size": int(len(pool)),
    }
    return picked.reset_index(drop=True), rule


def parse_pager_xml(text):
    root = ET.fromstring(text)
    bins = {}
    for node in root.findall("exposure"):
        if "dmin" not in node.attrib:
            continue
        mmi = int(round(float(node.attrib["dmin"]) + 0.5))
        bins[mmi] = int(float(node.attrib.get("exposure", 0) or 0))
    event = root.find("event")
    # newer files put <alert> under the root, older ones nest them
    alerts = {a.attrib.get("type"): a.attrib.get("level") for a in root.iter("alert")}
    out = {f"pop_mmi{k}": bins.get(k, 0) for k in MMI_BINS}
    out.update({
        "pager_ccode": root.attrib.get("ccode"),
        "local_time": event.attrib.get("localtime") if event is not None else None,
        "max_mmi": float(event.attrib["maxmmi"]) if event is not None and event.attrib.get("maxmmi") else None,
        "fatality_alert": alerts.get("fatality"),
        "economic_alert": alerts.get("economic"),
        "pager_version": event.attrib.get("number") if event is not None else None,
    })
    return out


def pager_url_for(event_id):
    r = fetch.get(COMCAT, params={"eventid": event_id, "format": "geojson"})
    if r.status_code != 200:
        return None
    products = r.json()["properties"].get("products", {})
    pager = products.get("losspager")
    if not pager:
        return None
    contents = pager[0].get("contents", {})
    return {
        "pager_xml": contents.get("pager.xml", {}).get("url"),
        "exposures_json": contents.get("json/exposures.json", {}).get("url"),
    }


def load_pager_cache():
    path = PAGER_DIR / "parsed.jsonl"
    done = {}
    if path.exists():
        for line in path.read_text().splitlines():
            try:
                row = json.loads(line)
            except json.JSONDecodeError:
                continue  # a line cut short by an interrupted run is fetched again
            done[row["event_id"]] = row
    return done


def fetch_pager_rows(event_ids, keep_raw):
    cache = load_pager_cache()
    # network failures are retried on the next run; parse failures and missing products are not
    todo = [e for e in event_ids if e not in cache or cache[e].get("error", "").startswith("fetch failed")]
    out_path = PAGER_DIR / "parsed.jsonl"
    write_lock = threading.Lock()
    print(f"pager.xml: {len(cache)} cached, {len(todo)} to fetch")

    def work(event_id):
        try:
            return fetch_one(event_id)
        except RuntimeError as err:
            row = {"event_id": event_id, "ok": False, "error": f"fetch failed: {err}"}
            with write_lock:
                with out_path.open("a") as fh:
                    fh.write(json.dumps(row) + "\n")
            return row

    def fetch_one(event_id):
        urls = pager_url_for(event_id)
        row = {"event_id": event_id, "ok": False}
        if urls and urls["pager_xml"]:
            r = fetch.get(urls["pager_xml"])
            if r.status_code == 200:
                try:
                    row.update(parse_pager_xml(r.content))
                    row["ok"] = True
                except ET.ParseError as err:
                    row["error"] = f"parse: {err}"
                if event_id in keep_raw:
                    (PAGER_DIR / "raw").mkdir(parents=True, exist_ok=True)
                    (PAGER_DIR / "raw" / f"{event_id}.pager.xml").write_bytes(r.content)
                    if urls["exposures_json"]:
                        e = fetch.get(urls["exposures_json"])
                        if e.status_code == 200:
                            (PAGER_DIR / "raw" / f"{event_id}.exposures.json").write_bytes(e.content)
            else:
                row["error"] = f"pager.xml http {r.status_code}"
        else:
            row["error"] = "no losspager pager.xml"
        with write_lock:
            with out_path.open("a") as fh:
                fh.write(json.dumps(row) + "\n")
        return row

    started = time.time()
    with ThreadPoolExecutor(max_workers=8) as pool:
        futures = [pool.submit(work, e) for e in todo]
        for i, fut in enumerate(as_completed(futures), 1):
            row = fut.result()
            cache[row["event_id"]] = row
            if i % 200 == 0:
                rate = i / (time.time() - started)
                print(f"  {i}/{len(todo)} ({rate:.1f} events/s)", flush=True)
    return cache


# --- NCEI HazEL -------------------------------------------------------------------------------

def load_ncei():
    pages = []
    page = 1
    while True:
        path = RAW / "ncei" / f"page_{page:03d}.json"
        cached_file(HAZEL, path, params={"itemsPerPage": 200, "page": page})
        data = json.loads(path.read_text())
        pages.extend(data["items"])
        if page >= data["totalPages"]:
            break
        page += 1
    df = pd.DataFrame(pages)
    has_time = df["year"].notna() & df["month"].notna() & df["day"].notna() & df["hour"].notna() & df["minute"].notna()
    df = df[has_time & (df["year"] >= 1900)].copy()
    sec = df["second"].fillna(0).astype(float)
    df["time"] = pd.to_datetime(dict(year=df["year"].astype(int), month=df["month"].astype(int),
                                     day=df["day"].astype(int), hour=df["hour"].astype(int),
                                     minute=df["minute"].astype(int)), utc=True, errors="coerce")
    df["time"] = df["time"] + pd.to_timedelta(sec, unit="s")
    df["ncei_mag"] = df["eqMagnitude"]
    for alt in ("eqMagMw", "eqMagMs", "eqMagMb", "eqMagMl"):
        if alt in df:
            df["ncei_mag"] = df["ncei_mag"].fillna(df[alt])
    return df.dropna(subset=["time"]).reset_index(drop=True)


def join_by_origin(left, right, right_cols, seconds=JOIN_SECONDS, km=JOIN_KM, dmag=JOIN_MAG,
                   right_mag="ncei_mag"):
    """Best right-row match for each left row: same origin within the time, distance and magnitude
    windows, nearest in time. A right row is used at most once."""
    r_times = epoch_seconds(right["time"])
    order = np.argsort(r_times)
    r_sorted = r_times[order]
    candidates = []
    for i, row in enumerate(left.itertuples(index=False)):
        t = epoch_seconds(row.time)
        lo = np.searchsorted(r_sorted, t - seconds, "left")
        hi = np.searchsorted(r_sorted, t + seconds, "right")
        for j in order[lo:hi]:
            rr = right.iloc[j]
            dist = float(haversine_km(row.lat, row.lon, rr["latitude"], rr["longitude"]))
            if dist > km:
                continue
            m = rr.get(right_mag)
            if pd.notna(m) and pd.notna(row.magnitude) and abs(float(m) - float(row.magnitude)) > dmag:
                continue
            candidates.append((abs(r_times[j] - t), i, j))
    candidates.sort()
    used_left, used_right, pairs = set(), set(), {}
    for _, i, j in candidates:
        if i in used_left or j in used_right:
            continue
        used_left.add(i)
        used_right.add(j)
        pairs[i] = j
    matched = pd.DataFrame(index=left.index, columns=right_cols, dtype=object)
    for i, j in pairs.items():
        matched.iloc[i] = right.iloc[j][right_cols].to_numpy()
    return matched, len(pairs)


def ncei_labels(matched):
    deaths = pd.to_numeric(matched["deaths"], errors="coerce")
    deaths = deaths.fillna(pd.to_numeric(matched["deathsTotal"], errors="coerce"))
    injured = pd.to_numeric(matched["injuries"], errors="coerce")
    injured = injured.fillna(pd.to_numeric(matched["injuriesTotal"], errors="coerce"))
    return deaths, injured


# --- EXPO-CAT / PAGER-CAT ---------------------------------------------------------------------

def expocat_bins(frame):
    """Integer MMI bin k in pager.xml covers [k-0.5, k+0.5). EXPO-CAT columns hold people at
    MMI x +/- 0.25 in 0.5 steps, so bin k = half of column k-0.5 + column k + half of column k+0.5.
    Bin 4 is short by the missing 3.5 column (below the catalogue's range)."""
    def col(prefix, mmi):
        name = f"{prefix}{int(round(mmi * 10)):03d}"
        return frame[name].fillna(0) if name in frame else 0.0

    out = {}
    for k in FEATURE_BINS:
        total = 0.0
        for prefix in ("U", "R"):
            total = total + 0.5 * col(prefix, k - 0.5) + col(prefix, k) + 0.5 * col(prefix, k + 0.5)
        out[f"pop_mmi{k}"] = total
    return pd.DataFrame(out)


def load_expocat():
    path = cached_file(EXPOCAT_URL, CATALOGS / "EXPO_CAT_2007_12.csv")
    df = pd.read_csv(path)
    zpath = CATALOGS / "PAGER_CAT_v2.zip"
    if not zpath.exists():
        listing = fetch.get("https://www.sciencebase.gov/catalog/item/5bc730dde4b0fc368ebcad8a",
                            params={"format": "json"}).json()
        url = next(f["url"] for f in listing.get("files", []) if f["name"].endswith(".zip"))
        cached_file(url, zpath)
    with zipfile.ZipFile(zpath) as zf:
        name = next(n for n in zf.namelist() if n.endswith(".csv"))
        with zf.open(name) as fh:
            pcat = pd.read_csv(fh, encoding="latin-1", low_memory=False,
                               usecols=["eqID", "PAGER_prefInjuries", "PAGER_prefInjuriesSrc"])
    # PAGER-CAT ids carry seconds (YYYYMMDDHHMMSS); EXPO-CAT ids stop at the minute
    pcat["eqID"] = pcat["eqID"] // 100
    same_minute = int(pcat["eqID"].duplicated(keep=False).sum())
    pcat = pcat.drop_duplicates("eqID", keep=False)
    df = df.merge(pcat, on="eqID", how="left")

    shaking = df["PAGER_prefShakingDeaths"]
    total = df["PAGER_prefTotalDeaths"]
    df["deaths"] = shaking
    df.loc[shaking.isna() & total.isna(), "deaths"] = 0.0
    dropped_secondary = int((shaking.isna() & total.notna()).sum())
    df = df[df["deaths"].notna()].copy()

    bins = expocat_bins(df)
    out = pd.DataFrame({
        "event_id": "expocat:" + df["eqID"].astype(str),
        "source": "expocat",
        "time": pd.to_datetime(df["PAGER_dateStr"], utc=True, errors="coerce"),
        "lat": df["PAGER_lat"], "lon": df["PAGER_lon"],
        "depth_km": pd.to_numeric(df["PAGER_depth"], errors="coerce"),
        "magnitude": df["PAGER_prefMag"],
        "local_hour": pd.to_numeric(df["localHour"], errors="coerce"),
        "iso2": df["ISO_code"].replace({"UK": np.nan}),
        "pager_alert": np.nan,
        "deaths": df["deaths"].astype(float),
        "injured": pd.to_numeric(df["PAGER_prefInjuries"], errors="coerce"),
        "weight": 1.0,
    })
    out = pd.concat([out.reset_index(drop=True), bins.reset_index(drop=True)], axis=1)
    return out.dropna(subset=["time"]), {
        "expocat_rows": int(len(shaking)),
        "dropped_only_total_deaths_known": dropped_secondary,
        "pagercat_join": "PAGER-CAT eqID // 100 == EXPO-CAT eqID; PAGER-CAT minutes holding two events "
                         f"are skipped ({same_minute} rows)",
        "injuries_known": int(out["injured"].notna().sum()),
        "deaths_rule": "PAGER_prefShakingDeaths; 0 when both shaking and total deaths are blank (no source "
                       "catalogue reported any); dropped when only total deaths are known",
    }


# --- countries, income, GDP -------------------------------------------------------------------

class Countries:
    def __init__(self):
        path = cached_file(NE_URL, CATALOGS / "ne_50m_admin_0_countries.geojson")
        feats = json.loads(path.read_text())["features"]
        self.shapes = []
        verts = []
        owners = []
        for idx, f in enumerate(feats):
            p = f["properties"]
            iso3 = p.get("ISO_A3_EH") if p.get("ISO_A3_EH") not in (None, "-99") else p.get("ADM0_A3")
            geom = f["geometry"]
            polys = geom["coordinates"] if geom["type"] == "MultiPolygon" else [geom["coordinates"]]
            for poly in polys:
                ring = np.asarray(poly[0])
                self.shapes.append((iso3, MplPath(ring)))
                verts.append(ring)
                owners.extend([iso3] * len(ring))
        self.verts = np.vstack(verts)
        self.owners = np.asarray(owners)

    def locate(self, lats, lons, max_km=500):
        pts = np.column_stack([lons, lats])
        found = np.full(len(pts), None, dtype=object)
        for iso3, path in self.shapes:
            empty = found == None  # noqa: E711
            if not empty.any():
                break
            inside = path.contains_points(pts[empty])
            idx = np.flatnonzero(empty)[inside]
            found[idx] = iso3
        for i in np.flatnonzero(found == None):  # noqa: E711
            d = haversine_km(lats[i], lons[i], self.verts[:, 1], self.verts[:, 0])
            j = int(np.argmin(d))
            if d[j] <= max_km:
                found[i] = self.owners[j]
        return found


def load_world_bank():
    meta_path = cached_file(WB_COUNTRIES, CATALOGS / "wb_countries.json")
    meta = json.loads(meta_path.read_text())[1]
    iso2_to_3 = {m["iso2Code"]: m["id"] for m in meta}
    region = {m["id"]: m["region"]["value"].strip() for m in meta if m["region"]["value"].strip() != "Aggregates"}
    current_income = {m["id"]: API_INCOME.get(m["incomeLevel"]["id"]) for m in meta}

    wb = openpyxl.load_workbook(cached_file(OGHIST_URL, CATALOGS / "OGHIST.xlsx"), read_only=True)
    rows = list(wb["Country Analytical History"].iter_rows(values_only=True))
    years = [v for v in rows[5][2:] if isinstance(v, int)]
    history = {}
    for r in rows[11:]:
        if not r[0] or not isinstance(r[0], str) or len(r[0]) != 3:
            continue
        vals = {}
        for y, v in zip(years, r[2:2 + len(years)]):
            if isinstance(v, str) and v.strip() in INCOME_ORDER:
                vals[y] = v.strip()
        if vals:
            history[r[0]] = vals

    gdp_path = cached_file(WB_GDP, CATALOGS / "wb_gdp_per_capita.json")
    gdp_rows = json.loads(gdp_path.read_text())[1]
    gdp = {}
    for g in gdp_rows:
        if g["value"] is not None:
            gdp.setdefault(g["countryiso3code"], {})[int(g["date"])] = float(g["value"])
    return iso2_to_3, region, current_income, history, gdp, (min(years), max(years))


def income_for(iso3, year, history, current_income, span):
    """Class in force during `year` ~ the classification made from the previous calendar year's
    data. Years outside the table use its nearest edge; countries missing from it use today's class."""
    if isinstance(iso3, str) and iso3 in history:
        vals = history[iso3]
        want = min(max(year - 1, span[0]), span[1])
        if want in vals:
            return vals[want], "oghist"
        nearest = min(vals, key=lambda y: abs(y - want))
        return vals[nearest], "oghist-nearest"
    if isinstance(iso3, str) and current_income.get(iso3):
        return current_income[iso3], "api-current"
    return None, "unknown"


def gdp_for(iso3, year, gdp):
    series = gdp.get(iso3) if isinstance(iso3, str) else None
    if not series:
        return np.nan
    want = year - 1
    nearest = min(series, key=lambda y: abs(y - want))
    if abs(nearest - want) > 10:
        return np.nan
    return series[nearest]


# --- sequences and holdouts -------------------------------------------------------------------

def gk_window(mag):
    """Gardner & Knopoff (1974) windows in the fitted form given by van Stiphout et al. (2012)."""
    dist_km = 10 ** (0.1238 * mag + 0.983)
    days = 10 ** (0.032 * mag + 2.7389) if mag >= 6.5 else 10 ** (0.5409 * mag - 0.547)
    return dist_km, days


def group_sequences(df):
    """Largest unassigned event opens a group; unassigned events inside its distance window and
    between min(window, 30 days) before and the full window after join it."""
    order = df["magnitude"].fillna(0).sort_values(ascending=False).index
    t = epoch_seconds(df["time"]) / 86400
    lat = df["lat"].to_numpy()
    lon = df["lon"].to_numpy()
    group = pd.Series(-1, index=df.index)
    pos = {idx: k for k, idx in enumerate(df.index)}
    g = 0
    for idx in order:
        if group[idx] != -1:
            continue
        k = pos[idx]
        dist_km, days = gk_window(float(df.at[idx, "magnitude"] or 0))
        dt = t - t[k]
        near = (dt >= -min(days, 30)) & (dt <= days)
        near &= haversine_km(lat[k], lon[k], lat, lon) <= dist_km
        near &= (group == -1).to_numpy()
        group[df.index[near]] = g
        group[idx] = g
        g += 1
    return group


def holdout_flags(df):
    flags = pd.Series("", index=df.index)
    t = df["time"]
    for name, h in HOLDOUTS.items():
        t0 = pd.Timestamp(h["time"])
        dt_days = (t - t0).dt.total_seconds() / 86400
        near = (dt_days >= -HOLDOUT_BEFORE_DAYS) & (dt_days <= HOLDOUT_AFTER_DAYS)
        near &= haversine_km(h["lat"], h["lon"], df["lat"].to_numpy(), df["lon"].to_numpy()) <= HOLDOUT_RADIUS_KM
        near |= df["event_id"].isin([h["id"], *h["extra_ids"]])
        flags[near] = name
    return flags


# --- main -------------------------------------------------------------------------------------

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--green-sample", type=int, default=1500)
    ap.add_argument("--seed", type=int, default=7)
    ap.add_argument("--refresh-year", type=int, default=None)
    ap.add_argument("--cached-only", action="store_true", help="skip events not downloaded yet (quick check)")
    args = ap.parse_args()
    t_start = time.time()
    summary = {"built_at": datetime.now(timezone.utc).isoformat(timespec="seconds")}

    events = list_pager_events(args.refresh_year)
    picked, rule = pick_events(events, args.green_sample, args.seed)
    summary["sampling"] = rule
    print(f"ComCat events with PAGER: {len(events)}, picked {len(picked)}")

    if args.cached_only:
        picked = picked[picked["event_id"].isin(load_pager_cache())].reset_index(drop=True)
    keep_raw = {h["id"] for h in HOLDOUTS.values()} | {"se609212"}
    pager = fetch_pager_rows(picked["event_id"].tolist(), keep_raw)
    parsed = pd.DataFrame([pager[e] for e in picked["event_id"] if e in pager])
    comcat = picked.merge(parsed, on="event_id", how="left")
    failed = comcat[comcat["ok"] != True]  # noqa: E712
    summary["comcat"] = {"picked": int(len(picked)), "pager_parse_failed": int(len(failed)),
                         "failure_reasons": failed["error"].value_counts().to_dict() if len(failed) else {}}
    comcat = comcat[comcat["ok"] == True].copy()  # noqa: E712
    # failed downloads are treated as random, so the surviving sampled rows still stand for the whole pool
    n_ok_sampled = int(comcat["sampled"].sum())
    if n_ok_sampled:
        comcat.loc[comcat["sampled"], "weight"] = rule["pool_size"] / n_ok_sampled
    # older files give "HH:MM:SS", newer ones a full date and time
    hm = comcat["local_time"].astype(str).str.extract(r"(\d{1,2}):(\d{2}):\d{2}\s*$").astype(float)
    comcat["local_hour"] = hm[0] + hm[1] / 60
    comcat["pager_alert"] = comcat["fatality_alert"].map({"green": 0, "yellow": 1, "orange": 2, "red": 3})
    comcat["source"] = "comcat"

    ncei = load_ncei()
    summary["ncei_rows_with_time"] = int(len(ncei))
    right_cols = ["id", "deaths", "deathsTotal", "injuries", "injuriesTotal", "country", "locationName"]
    for c in right_cols:
        if c not in ncei:
            ncei[c] = np.nan
    matched, n_matched = join_by_origin(comcat.reset_index(drop=True), ncei, right_cols)
    comcat = comcat.reset_index(drop=True)
    deaths, injured = ncei_labels(matched)
    comcat["ncei_id"] = matched["id"]
    comcat["ncei_country"] = matched["country"]
    comcat["deaths"] = deaths.fillna(0.0).astype(float)
    comcat["injured"] = injured.astype(float)
    summary["ncei_join"] = {
        "rule": f"origin time within {JOIN_SECONDS} s, epicentres within {JOIN_KM} km, magnitude within "
                f"{JOIN_MAG}; nearest in time wins; each NCEI row used once",
        "comcat_matched": int(n_matched),
        "deaths_label": "NCEI `deaths`, else `deathsTotal`; unmatched events = 0 (NCEI lists every "
                        "earthquake that caused deaths)",
        "injured_label": "NCEI `injuries`, else `injuriesTotal`; unknown otherwise (row excluded from the "
                         "injury model)",
    }

    expo, expo_info = load_expocat()
    summary["expocat"] = expo_info
    # EXPO-CAT rows that ComCat also covers are dropped in favour of the ComCat row
    dup, n_dup = join_by_origin(
        expo.reset_index(drop=True),
        comcat.rename(columns={"lat": "latitude", "lon": "longitude", "magnitude": "cc_mag"}),
        ["event_id"], right_mag="cc_mag")
    expo = expo.reset_index(drop=True)[dup["event_id"].isna().to_numpy()]
    summary["expocat"]["dropped_as_comcat_duplicates"] = int(n_dup)

    cols = ["event_id", "source", "time", "lat", "lon", "depth_km", "magnitude", "local_hour", "pager_alert",
            "deaths", "injured", "weight"] + [f"pop_mmi{k}" for k in FEATURE_BINS]
    table = pd.concat([comcat[cols + ["pager_ccode", "summary_alert", "max_mmi", "ncei_id"]],
                       expo[cols + ["iso2"]]], ignore_index=True)

    countries = Countries()
    iso2_to_3, region, current_income, history, gdp, span = load_world_bank()
    iso3 = pd.Series(countries.locate(table["lat"].to_numpy(), table["lon"].to_numpy()), index=table.index)
    from_expo = table["iso2"].map(iso2_to_3)
    iso3 = from_expo.where(from_expo.notna(), iso3)
    ccode3 = table.get("pager_ccode").map(iso2_to_3) if "pager_ccode" in table else None
    if ccode3 is not None:
        iso3 = iso3.where(iso3.notna(), ccode3)
    table["iso3"] = iso3
    years = table["time"].dt.year
    inc = [income_for(c, int(y), history, current_income, span) for c, y in zip(table["iso3"], years)]
    table["income_class"] = [INCOME_ORDER.get(v) if v else np.nan for v, _ in inc]
    table["income_source"] = [s for _, s in inc]
    table["gdp_per_capita"] = [gdp_for(c, int(y), gdp) for c, y in zip(table["iso3"], years)]
    table["region"] = table["iso3"].map(region).fillna("Unknown")
    summary["country"] = {
        "rule": "EXPO-CAT ISO code where given; otherwise the Natural Earth 1:50m country containing the "
                "epicentre, or the nearest one within 500 km; otherwise PAGER's ccode",
        "unknown_country_rows": int(table["iso3"].isna().sum()),
        "income_sources": table["income_source"].value_counts().to_dict(),
        "income_rule": "World Bank OGHIST class for calendar-year data (event year - 1), clipped to "
                       f"{span[0]}-{span[1]}; nearest listed year if that one is blank; countries not in "
                       "OGHIST use today's API class",
        "gdp_rule": "World Bank NY.GDP.PCAP.CD for (event year - 1), nearest year within 10 years",
    }

    table["sequence"] = group_sequences(table)
    table["holdout"] = holdout_flags(table)
    held_groups = set(table.loc[table["holdout"] != "", "sequence"])
    table.loc[table["sequence"].isin(held_groups) & (table["holdout"] == ""), "holdout"] = "sequence-of-holdout"
    summary["sequences"] = {
        "rule": "Gardner & Knopoff (1974) space-time windows (van Stiphout et al. 2012 fit), largest event "
                "first; foreshock side capped at 30 days",
        "groups": int(table["sequence"].nunique()),
    }
    summary["holdout_rule"] = (f"events within {HOLDOUT_RADIUS_KM} km and {HOLDOUT_BEFORE_DAYS} days before to "
                               f"{HOLDOUT_AFTER_DAYS} days after each holdout mainshock, the listed ids, and "
                               "anything grouped into the same sequence")

    table = table.drop(columns=["iso2", "pager_ccode"], errors="ignore")
    table["time"] = table["time"].dt.strftime("%Y-%m-%dT%H:%M:%SZ")
    ARTIFACTS.mkdir(parents=True, exist_ok=True)
    table.to_csv(ARTIFACTS / "training_table.csv.gz", index=False)

    train = table[table["holdout"] == ""]
    summary["rows"] = {
        "total": int(len(table)),
        "by_source": table["source"].value_counts().to_dict(),
        "holdout": table["holdout"].value_counts().to_dict(),
        "train_deaths_model": int(len(train)),
        "train_deaths_gt0": int((train["deaths"] > 0).sum()),
        "train_injury_model": int(train["injured"].notna().sum()),
        "train_injured_gt0": int((train["injured"] > 0).sum()),
    }
    summary["http_requests_this_run"] = fetch.calls
    summary["wall_seconds"] = round(time.time() - t_start, 1)
    (ARTIFACTS / "dataset_summary.json").write_text(json.dumps(summary, indent=2, default=str))
    print(json.dumps(summary["rows"], indent=2))


if __name__ == "__main__":
    main()
